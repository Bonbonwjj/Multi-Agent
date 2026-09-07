from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "ChatDev_ReConcile_Group_Agent_架构说明.docx"
NAVY = "17365D"
PALE = "EAF0F7"
LIGHT = "F5F7FA"
BORDER = "D9D9D9"
CHINESE_FONT = "Heiti SC"


def shade(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    element = tbl_pr.first_child_found_in("w:tblBorders")
    if element is None:
        element = OxmlElement("w:tblBorders")
        tbl_pr.append(element)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "5")
        tag.set(qn("w:color"), BORDER)
        element.append(tag)


def set_cell_margin(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)


def table(doc, headers, rows, widths=None):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.autofit = False
    tbl.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for index, header in enumerate(headers):
        cell = tbl.rows[0].cells[index]
        cell.text = header
        shade(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9.5)
    for row_index, row in enumerate(rows):
        cells = tbl.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = str(value)
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                shade(cells[index], LIGHT)
            for run in cells[index].paragraphs[0].runs:
                run.font.size = Pt(9.5)
    if widths:
        for row in tbl.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    for row in tbl.rows:
        for cell in row.cells:
            set_cell_margin(cell)
    borders(tbl)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return tbl


def code(doc, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.right_indent = Inches(0.25)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.05
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F1F3F5")
    pPr.append(shd)
    run = p.add_run(text)
    run.font.name = "Courier New"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), CHINESE_FONT)
    run.font.size = Pt(8.5)


def apply_chinese_font(doc) -> None:
    """Set an explicit CJK font on every run for reliable Word/PDF rendering."""
    paragraphs = list(doc.paragraphs)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for section in doc.sections:
        paragraphs.extend(section.header.paragraphs)
        paragraphs.extend(section.footer.paragraphs)
    for paragraph in paragraphs:
        for run in paragraph.runs:
            run.font.name = CHINESE_FONT
            fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
            for key in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                fonts.set(qn(key), CHINESE_FONT)


def heading(doc, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def bullet(doc, text: str, level: int = 0) -> None:
    style = "List Bullet" if level == 0 else "List Bullet 2"
    doc.add_paragraph(text, style=style)


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), CHINESE_FONT)
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.15
    for name, size in (("Title", 25), ("Heading 1", 17), ("Heading 2", 13), ("Heading 3", 11)):
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), CHINESE_FONT)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
        if name == "Title":
            style_ppr = style._element.get_or_add_pPr()
            style_border = style_ppr.first_child_found_in("w:pBdr")
            if style_border is not None:
                style_ppr.remove(style_border)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(75)
    title.add_run("ChatDev ReConcile Group Agent 架构说明")
    title_ppr = title._p.get_or_add_pPr()
    title_border = title_ppr.first_child_found_in("w:pBdr")
    if title_border is not None:
        title_ppr.remove(title_border)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("从单角色智能体扩展到可组合的多专家协作框架").bold = True
    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    intro.paragraph_format.space_before = Pt(24)
    intro.add_run("项目路径\n").bold = True
    intro.add_run(str(ROOT))
    doc.add_paragraph("版本 1.0    2026 年 9 月", style=None).alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    heading(doc, "文档目的")
    doc.add_paragraph(
        "本文档说明如何在 ChatDev 的阶段式开发流程上增加 Group Agent 抽象，并以 ReConcile 作为组内协作算法。"
        "读者可以据此理解现有代码、创建具有独立 Character 和 Skills 的专家、把 CEO 或 CTO 等角色替换为专家组，"
        "并运行可追踪的共识形成过程。当前实现保留原有单 Agent 接口，因此扩展不会破坏既有八阶段流程。"
    )
    heading(doc, "核心结论", 2)
    bullet(doc, "Character 定义专家的目标、性格、推理风格和分歧处理方式。")
    bullet(doc, "Skill 定义领域知识、工作步骤和结果校验规则。")
    bullet(doc, "GroupAgent 将 Character、多个 Skills 与 LLM 组合为一个专家。")
    bullet(doc, "AgentGroup 使用可替换的 CollaborationStrategy 协调多个专家。")
    bullet(doc, "GroupRoleAgent 使专家组能够直接替换 ChatDev 的 CEO、CTO、Programmer 等角色。")
    bullet(doc, "ReConcile 先独立回答，再共享证据并重新判断，最后执行证据加权聚合。")

    heading(doc, "系统总体架构")
    code(doc, "User Task\n  ↓\nChatDevPipeline\n  ↓ role_executors[role]\nGroupRoleAgent\n  ↓\nAgentGroup\n  ↓ ReConcileStrategy\nGroupAgent A | GroupAgent B | GroupAgent C\n  ↓ Character + Skills + LLM\nGroupResponse")
    doc.add_paragraph(
        "ChatDevPipeline 仍控制需求分析、语言选择、编码、补全、审查、测试和文档阶段。每个阶段通过角色名获取执行器。"
        "没有配置 Group 时，执行器仍是原来的 Agent；配置 GroupRoleAgent 后，相同的角色槽由多个专家共同完成。"
    )
    table(doc, ["层次", "主要对象", "职责"], [
        ("流程层", "ChatDevPipeline", "控制阶段顺序、记忆、代码产物和测试反馈"),
        ("兼容层", "GroupRoleAgent", "将 GroupResponse 转换为原 Agent 所需的文本回答"),
        ("群组层", "AgentGroup", "保存专家成员、领域和协作策略"),
        ("协作层", "ReConcileStrategy", "控制独立回答、重新判断、共识检测和停止"),
        ("专家层", "GroupAgent", "组合 Character、Skills 和 LLM"),
        ("能力层", "Skill", "注入专业知识、工作流和校验规则"),
    ], [0.8, 1.55, 3.9])

    heading(doc, "现有 ChatDev 代码导读")
    table(doc, ["文件", "作用", "扩展关系"], [
        ("agents.py", "原单角色 Agent、seminar、CDH", "新增 AgentLike 协议以接收单 Agent 或 Group"),
        ("pipeline.py", "官方八阶段 Chat Chain", "role_executors 支持按角色注入 Group"),
        ("llm.py", "OpenAI 兼容接口与离线 Mock", "每个 GroupAgent 可以拥有自己的 LLM"),
        ("skills.py", "Skill 协议与注册表", "领域知识扩展入口"),
        ("group_agents.py", "Character 和 GroupAgent", "单个专家的标准输入输出"),
        ("reconcile.py", "AgentGroup 与 ReConcile", "组内协作、聚合和 ChatDev 适配"),
        ("events.py", "ChatDev 事件记录", "后续可扩展 Group 轮次与奖励事件"),
    ], [1.15, 2.2, 3.0])

    heading(doc, "Character 接口")
    doc.add_paragraph("Character 只描述如何思考和交流，不直接承载专业知识。这样同一专业 Skill 可以被不同风格的专家复用。")
    code(doc, "@dataclass(frozen=True)\nclass Character:\n    id: str\n    role: str\n    goal: str\n    traits: tuple[str, ...]\n    reasoning_style: str\n    communication_style: str\n    disagreement_policy: str")
    table(doc, ["字段", "意义", "示例"], [
        ("role", "专家身份", "Security Architect"),
        ("goal", "当前优化目标", "最小化安全风险"),
        ("traits", "稳定行为倾向", "谨慎、怀疑、重视证据"),
        ("reasoning_style", "分析方式", "falsification"),
        ("communication_style", "表达方式", "concise"),
        ("disagreement_policy", "遇到分歧时的规则", "要求可验证证据"),
    ], [1.25, 2.25, 2.85])

    heading(doc, "Skill 接口")
    doc.add_paragraph(
        "Skill 是稳定的领域能力扩展点。instructions 方法向 Agent 注入知识和操作步骤；validate 方法对最终答案执行确定性校验。"
        "框架目前提供 PromptSkill、MarkdownSkill 和 SkillRegistry，并允许未来增加工具型 Skill。"
    )
    code(doc, "class Skill(Protocol):\n    id: str\n    description: str\n\n    def instructions(self, context: SkillContext) -> str: ...\n    def validate(self, answer: str) -> list[str]: ...")
    heading(doc, "Skill 使用示例", 2)
    code(doc, "product_skill = PromptSkill(\n    id=\"product_analysis\",\n    description=\"将需求转成产品定义\",\n    knowledge=\"优先满足硬约束和最小可行范围\",\n    workflow=(\"提取约束\", \"比较形态\", \"定义验收标准\"),\n    required_terms=(\"application\",),\n)")
    doc.add_paragraph(
        "required_terms 只是最小示例。真实学科 Skill 可以检查引用格式、必需章节、单位、风险声明或结构化字段。"
        "Skill 不直接选择最终答案，因而与 ReConcile、Debate 或未来的强化学习调度器保持解耦。"
    )

    heading(doc, "GroupAgent 接口")
    code(doc, "@dataclass\nclass GroupAgent:\n    id: str\n    character: Character\n    skills: list[Skill]\n    llm: LLM\n\n    def act(self, request: AgentRequest) -> AgentResponse: ...")
    table(doc, ["输入字段", "内容"], [
        ("task", "当前用户任务"),
        ("phase", "当前 ChatDev 阶段或领域步骤"),
        ("context", "上游产物、约束与外部工具结果"),
        ("peer_responses", "ReConcile 后续轮次中可见的同伴观点"),
        ("round_number", "独立回答为 0，重新判断从 1 开始"),
    ], [1.65, 4.7])
    table(doc, ["输出字段", "用途"], [
        ("answer", "候选结论"), ("claims", "可单独核验的主张"), ("evidence", "支持结论的依据"),
        ("uncertainties", "未解决的不确定性"), ("confidence", "0 到 1 的自评置信度"),
        ("skill_errors", "未满足 Skill 校验规则的项目"),
    ], [1.65, 4.7])

    heading(doc, "ReConcile 协作算法")
    code(doc, "Round 0  独立回答 隐藏同伴观点\n   ↓ 计算完全一致答案的比例\nRound 1  查看同伴答案 证据和置信度后重新判断\n   ↓ 未达到阈值则继续\nRound N  达成阈值或达到最大轮数\n   ↓\nEvidenceWeightedAggregator\n   ↓\n统一答案 赞成者 分歧 证据 置信度 完整轨迹")
    heading(doc, "共识检测", 2)
    doc.add_paragraph(
        "当前实现对规范化后的答案做精确分桶，一致率等于最大答案桶的成员数除以总成员数。该方法简单、可测试，"
        "适合框架第一版。未来可增加语义聚类器，处理措辞不同但含义相同的答案。"
    )
    heading(doc, "证据加权聚合", 2)
    code(doc, "score(response) = confidence\n                + evidence_bonus × min(3, evidence_count)\n                - skill_error_penalty × skill_error_count")
    doc.add_paragraph(
        "聚合器先按答案分组，再累加组内成员得分。它不把多数意见直接当作事实，也会保留所有少数意见。"
        "外部工具证据尚未单独分级，后续应给编译器、测试、权威数据库等可验证证据更高权重。"
    )

    heading(doc, "CEO Group 示例")
    table(doc, ["成员", "Character 目标", "共同 Skill"], [
        ("Product Strategist", "最大化产品价值", "product_analysis"),
        ("User Advocate", "保护可用性与用户需求", "product_analysis"),
        ("Risk Controller", "降低交付与范围风险", "product_analysis"),
    ], [1.7, 2.65, 2.0])
    doc.add_paragraph(
        "离线示例中，第一轮有两名专家选择 CLI application，一名专家选择 Desktop application。第一轮一致率为 2/3，"
        "低于默认 0.75，因此进入第二轮。User Advocate 阅读“零第三方依赖”和 Python 标准库证据后修改判断，最终三人一致。"
    )
    code(doc, "Consensus: CLI application\nConfidence: 0.92\nRounds: 2\nAgreement: strategist, user_advocate, risk_controller\nEvidence: Python standard library provides argparse and json")

    heading(doc, "接入原 ChatDev")
    code(doc, "ceo_group = AgentGroup(\n    id=\"ceo_group\",\n    domain=\"product\",\n    agents=[strategist, user_advocate, risk_controller],\n    strategy=ReConcileStrategy(max_rounds=3),\n)\n\npipeline = ChatDevPipeline(\n    llm=default_llm,\n    role_executors={\n        \"ceo\": GroupRoleAgent(\"ceo\", ceo_group),\n    },\n)")
    doc.add_paragraph(
        "同一方式可以替换 cto、programmer、reviewer 和 tester。不同 Group 的成员可以使用不同模型、Character 和 Skills；"
        "ChatDevPipeline 仍只依赖 AgentLike.respond，因此阶段逻辑无需了解组内实现。"
    )

    heading(doc, "运行指南")
    heading(doc, "执行离线 ReConcile 示例", 2)
    code(doc, "cd /Users/wjj/Documents/multi-agent/chatdev-paper-reproduction\nmake demo")
    heading(doc, "执行测试和源码编译检查", 2)
    code(doc, "make test")
    heading(doc, "执行原 ChatDev 离线示例", 2)
    code(doc, "make chatdev-demo")
    heading(doc, "执行真实模型示例", 2)
    code(doc, "export OPENAI_API_KEY=...\nexport OPENAI_BASE_URL=...\nexport OPENAI_MODEL=...\nPYTHONPATH=src python3 examples/reconcile_ceo_group.py")

    heading(doc, "测试覆盖")
    table(doc, ["测试", "验证内容"], [
        ("test_reconcile_reaches_consensus", "不同初始意见在证据共享后形成共识"),
        ("test_skill_registry_and_validation", "Skill 注册、解析和答案校验"),
        ("test_group_can_replace_legacy_role", "AgentGroup 通过适配器替换 CEO"),
        ("test_mock_pipeline", "原 ChatDev 八阶段仍能生成软件产物"),
        ("test_generated_paths", "模型生成文件不能逃逸输出目录"),
    ], [2.55, 3.8])

    heading(doc, "后续扩展路线")
    table(doc, ["优先级", "扩展", "建议"], [
        ("P0", "异步并行", "使用 async gather 并行执行独立回答，减少 Group 延迟"),
        ("P0", "结构化模型输出", "为 OpenAI 兼容接口增加 JSON Schema 或重试修复"),
        ("P1", "语义共识", "按主张和语义相似度聚类，不只比较完整字符串"),
        ("P1", "证据等级", "区分模型陈述、文档引用、工具结果和真实测试"),
        ("P1", "Group 事件流", "记录成员、轮次、观点变化、token 和延迟"),
        ("P2", "强化学习", "先学习路由、讨论停止和聚合权重，再考虑底层模型"),
        ("P2", "学科 Skills", "按统计、医学、经济等领域建立专家与验证规则"),
    ], [0.65, 1.45, 4.25])

    heading(doc, "设计限制与安全边界")
    bullet(doc, "相同底层模型的多个 Agent 并不等于真正独立的人类专家，可能共享同一知识盲区。")
    bullet(doc, "Agent 的自报置信度未必校准，最终判断应优先依赖外部证据和测试结果。")
    bullet(doc, "共识不代表正确；少数意见必须保留，不能为了达到阈值强制统一。")
    bullet(doc, "模型生成代码默认只做静态检查，只有在可信沙箱中才应启用执行。")
    bullet(doc, "专业 Skill 应明确适用范围、证据标准、限制和升级到人工专家的条件。")

    heading(doc, "验收结论")
    doc.add_paragraph(
        "当前版本已经形成可运行的第一阶段框架：原 ChatDev 流程保持兼容；CEO 等角色可替换为 AgentGroup；"
        "GroupAgent 具有独立 Character 和多个 Skills；ReConcile 能执行独立回答、同伴讨论、共识检测和证据加权聚合；"
        "离线示例与自动化测试可验证整个链路。下一阶段应优先增加异步并行、结构化输出约束和细粒度事件记录。"
    )

    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run("ChatDev ReConcile Group Agent 架构说明")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    apply_chinese_font(doc)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
